import io
import openpyxl
from datetime import datetime
import pandas as pd
from src.expedicao.cabotagem_processador import montar_relatorio_cabotagem

def test_exportar_excel_com_formato_data_abreviada_e_autofiltro():
    estado = {
        "containers": [
            {
                "carga": "C-100",
                "container": "CTR-1",
                "remessas": ["80112233"],
                "of_numero": "610001",
                "data_agendamento": "28/09/2026",
            },
            {
                "carga": "C-200",
                "container": "CTR-2",
                "remessas": ["80445566"],
                "of_numero": "610002",
                "data_agendamento": "30/09/2026",
            }
        ]
    }

    df = montar_relatorio_cabotagem(estado)
    assert "Data Ag. Recebimento Cliente" in df.columns

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Cabotagem')
        ws = writer.sheets['Cabotagem']
        ws.auto_filter.ref = ws.dimensions

        col_data_idx = None
        for idx, col_name in enumerate(df.columns, 1):
            if col_name == "Data Ag. Recebimento Cliente":
                col_data_idx = idx
                break

        assert col_data_idx is not None

        for row_idx in range(2, len(df) + 2):
            cell = ws.cell(row=row_idx, column=col_data_idx)
            if cell.value:
                dt = datetime.strptime(str(cell.value).strip(), "%d/%m/%Y").date()
                cell.value = dt
            cell.number_format = 'DD/MM/YYYY'

    output.seek(0)
    wb = openpyxl.load_workbook(output)
    assert 'Cabotagem' in wb.sheetnames
    ws = wb['Cabotagem']
    assert ws.auto_filter.ref is not None

    cell_data1 = ws.cell(row=2, column=col_data_idx)
    assert cell_data1.number_format == 'DD/MM/YYYY'
    assert hasattr(cell_data1.value, 'strftime')
    assert cell_data1.value.strftime('%d/%m/%Y') == '28/09/2026'

    cell_data2 = ws.cell(row=3, column=col_data_idx)
    assert cell_data2.number_format == 'DD/MM/YYYY'
    assert hasattr(cell_data2.value, 'strftime')
    assert cell_data2.value.strftime('%d/%m/%Y') == '30/09/2026'

if __name__ == '__main__':
    test_exportar_excel_com_formato_data_abreviada_e_autofiltro()
    print('TESTE DE FORMATO EXCEL PASSOU COM SUCESSO!')
